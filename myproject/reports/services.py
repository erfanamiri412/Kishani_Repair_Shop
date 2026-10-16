import csv
import hashlib
from io import BytesIO, StringIO
from pathlib import Path

from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone

from openpyxl import Workbook

from .analytics import AnalyticsService
from .charts import ChartsService
from .financial_reports import FinancialReportService
from .inventory_reports import InventoryReportService
from .invoicing_reports import InvoiceReportsService
from .employee_reports import EmployeeReportService
from .ml_productions import MLProductionService
from .monthly_reports import MonthlyReportService
from .yearly_reports import YearlyReportService
from .pdf_generator import ManagementReportPDFGenerator
from .models import (
    BackupHistory,
    ReportHistory,
)


class ReportsService:

    """
    سرویس مرکزی گزارش‌گیری تعمیرگاه

    داده گزارش ابتدا از دیتابیس جمع‌آوری می‌شود،
    سپس تحلیل و نمودار روی همان نسخه انجام می‌شود
    و در پایان نسخه نهایی گزارش به صورت ثابت ذخیره می‌شود.
    """

    @staticmethod
    def generate_financial_report(
        start_date=None,
        end_date=None,
    ):

        return FinancialReportService.financial_dashboard(
            start_date=start_date,
            end_date=end_date,
        )

    @staticmethod
    def generate_inventory_report():

        return InventoryReportService.inventory_summary()

    @staticmethod
    def generate_invoice_report():

        return InvoiceReportsService.financial_summary()

    @staticmethod
    def generate_employee_report():

        return EmployeeReportService.get_employee_ranking()

    @staticmethod
    def generate_analytics_report(
        start_date,
        end_date,
    ):

        return AnalyticsService.get_report_metrics(
            start_date,
            end_date,
        )

    @staticmethod
    def generate_charts(
        start_date,
        end_date,
    ):

        return ChartsService.period_charts(
            start_date,
            end_date,
        )

    @staticmethod
    def generate_monthly_report(
        year,
        month,
    ):

        return MonthlyReportService.generate(
            year=year,
            month=month,
        )

    @staticmethod
    def generate_yearly_report(
        year,
    ):

        return YearlyReportService.generate(
            year=year,
        )

    @staticmethod
    def generate_complete_report(
        start_date=None,
        end_date=None,
    ):

        today = timezone.localdate()
        if start_date is None and end_date is None:
            start_date = end_date = today
        elif start_date is None:
            start_date = end_date
        elif end_date is None:
            end_date = start_date

        report_data = ReportsService.generate_analytics_report(
            start_date,
            end_date,
        )
        report_data["charts"] = ChartsService.period_charts(start_date, end_date)
        report_data["period_start"] = start_date
        report_data["period_end"] = end_date
        report_data["report_type"] = "custom"
        report_data["month_label"] = "گزارش سفارشی"
        return report_data

    @staticmethod
    def generate_dashboard_data():

        return ReportsService.generate_complete_report()

    @staticmethod
    def export_monthly_pdf(
        year,
        month,
    ):

        report = ReportsService.generate_monthly_report(
            year=year,
            month=month,
        )

        return ManagementReportPDFGenerator().generate_monthly_report(
            report
        )

    @staticmethod
    def export_yearly_pdf(
        year,
    ):

        report = ReportsService.generate_yearly_report(
            year=year,
        )

        return ManagementReportPDFGenerator().generate_yearly_report(
            report
        )

    @staticmethod
    def _make_json_safe(value):

        if value is None:
            return None

        if hasattr(value, "isoformat"):
            return value.isoformat()

        if isinstance(value, dict):
            return {
                str(key): ReportsService._make_json_safe(item)
                for key, item in value.items()
            }

        if isinstance(value, (list, tuple)):
            return [
                ReportsService._make_json_safe(item)
                for item in value
            ]

        try:
            import decimal
            if isinstance(value, decimal.Decimal):
                if value == value.to_integral_value():
                    return int(value)
                return float(value)
        except Exception:
            pass

        if isinstance(value, (int, float, str, bool)):
            return value

        return str(value)

    @staticmethod
    def _get_report_path(report):

        if report.file_path:
            return report.file_path

        return None

    @staticmethod
    def _get_report_file_object(report):

        if not report.file_path:
            raise FileNotFoundError

        if report.file_path.startswith("/media/"):
            media_url = getattr(
                __import__("django.conf", fromlist=["settings"]).settings,
                "MEDIA_URL",
                "/media/",
            )
            file_name = report.file_path.replace(
                media_url,
                "",
                1,
            ).lstrip("/")
        else:
            file_name = report.file_path

        return default_storage.open(
            file_name,
            "rb",
        )

    @staticmethod
    def get_report_file_url(report):

        if not report.file_path:
            return None

        if report.file_path.startswith("/media/"):
            return report.file_path

        try:
            return default_storage.url(
                report.file_path
            )
        except Exception:
            return report.file_path

    @staticmethod
    def _save_pdf(
        report_type,
        year,
        month,
        pdf_data,
    ):

        if report_type == "monthly":
            filename = (
                f"reports/monthly/"
                f"{year}/{month:02d}/"
                f"monthly_{year}_{month:02d}.pdf"
            )
        else:
            filename = (
                f"reports/yearly/"
                f"{year}/"
                f"yearly_{year}.pdf"
            )

        # Regeneration after clearing ReportHistory must refresh the canonical PDF,
        # not create a suffixed duplicate such as monthly_1405_06_ab12cd.pdf.
        if default_storage.exists(filename):
            default_storage.delete(filename)
        saved_path = default_storage.save(
            filename,
            ContentFile(pdf_data),
        )

        return saved_path

    @staticmethod
    @transaction.atomic
    def create_monthly_report(
        year,
        month,
        user=None,
    ):

        generation_key = (
            f"monthly-{year}-{month:02d}"
        )

        existing = ReportHistory.objects.filter(
            generation_key=generation_key,
        ).first()

        if existing:
            return existing

        report_data = ReportsService.generate_monthly_report(
            year=year,
            month=month,
        )

        pdf_data = ManagementReportPDFGenerator().generate_monthly_report(
            report_data
        )

        file_path = ReportsService._save_pdf(
            "monthly",
            year,
            month,
            pdf_data,
        )

        snapshot = ReportsService._make_json_safe(
            report_data
        )

        file_hash = hashlib.sha256(
            pdf_data
        ).hexdigest()[:16]

        report = ReportHistory.objects.create(

            title=(
                f"گزارش جامع ماهانه - "
                f"{report_data['month_label']}"
            ),

            report_type="monthly",

            report_format="pdf",

            file_name=Path(
                file_path
            ).name,

            file_path=file_path,

            file_size=(
                f"{len(pdf_data) / 1024:.1f} KB"
            ),

            description=(
                "گزارش ماهانه قفل‌شده بر اساس داده‌های "
                "ثبت‌شده تا پایان دوره. "
                f"شناسه فایل: {file_hash}"
            ),

            period_year=year,

            period_month=month,

            period_start=report_data[
                "period_start"
            ],

            period_end=report_data[
                "period_end"
            ],

            generation_key=generation_key,

            snapshot=snapshot,

            chart_data=ReportsService._make_json_safe(
                report_data.get(
                    "charts",
                    {}
                )
            ),

            created_by=user,
        )

        return report

    @staticmethod
    @transaction.atomic
    def create_yearly_report(
        year,
        user=None,
    ):

        generation_key = (
            f"yearly-{year}"
        )

        existing = ReportHistory.objects.filter(
            generation_key=generation_key,
        ).first()

        if existing:
            return existing

        report_data = ReportsService.generate_yearly_report(
            year=year,
        )

        pdf_data = ManagementReportPDFGenerator().generate_yearly_report(
            report_data
        )

        file_path = ReportsService._save_pdf(
            "yearly",
            year,
            None,
            pdf_data,
        )

        snapshot = ReportsService._make_json_safe(
            report_data
        )

        file_hash = hashlib.sha256(
            pdf_data
        ).hexdigest()[:16]

        report = ReportHistory.objects.create(

            title=(
                f"گزارش جامع سالانه - {year}"
            ),

            report_type="yearly",

            report_format="pdf",

            file_name=Path(
                file_path
            ).name,

            file_path=file_path,

            file_size=(
                f"{len(pdf_data) / 1024:.1f} KB"
            ),

            description=(
                "گزارش سالانه قفل‌شده بر اساس داده‌های "
                "ثبت‌شده در کل سال. "
                f"شناسه فایل: {file_hash}"
            ),

            period_year=year,

            period_start=report_data[
                "period_start"
            ],

            period_end=report_data[
                "period_end"
            ],

            generation_key=generation_key,

            snapshot=snapshot,

            chart_data=ReportsService._make_json_safe(
                report_data.get(
                    "charts",
                    {}
                )
            ),

            created_by=user,
        )

        return report

    @staticmethod
    def create_previous_month_report(user=None):

        from .periods import (
            previous_jalali_month,
            today_jalali,
        )

        year, month, _ = today_jalali()

        year, month = previous_jalali_month(
            year,
            month,
        )

        return ReportsService.create_monthly_report(
            year=year,
            month=month,
            user=user,
        )

    @staticmethod
    def create_previous_year_report(user=None):

        from .periods import today_jalali

        year, _, _ = today_jalali()

        return ReportsService.create_yearly_report(
            year=year - 1,
            user=user,
        )

    @staticmethod
    def generate_scheduled_reports(
        monthly=True,
        yearly=True,
    ):

        """
        تولید idempotent گزارش‌های دوره‌های کامل‌شده.

        Scheduler می‌تواند روزانه اجرا شود.
        فقط دوره‌های کامل‌شده بررسی می‌شوند و گزارش‌های
        موجود دوباره ساخته نمی‌شوند.
        """

        from .periods import (
            previous_jalali_month,
            today_jalali,
        )

        year, month, _ = today_jalali()

        reports = []

        # ============================================================
        # MONTHLY REPORTS
        # ============================================================

        if monthly:

            # آخرین ماه کاملاً تکمیل‌شده
            last_year, last_month = previous_jalali_month(
                year,
                month,
            )

            # قدیمی‌ترین گزارش ماهانه موجود
            first_monthly = (
                ReportHistory.objects
                .filter(
                    report_type="monthly",
                    period_year__isnull=False,
                    period_month__isnull=False,
                )
                .order_by(
                    "period_year",
                    "period_month",
                )
                .first()
            )

            if first_monthly:
                start_year = first_monthly.period_year
                start_month = first_monthly.period_month
            else:
                start_year = last_year
                start_month = last_month

            current_year = start_year
            current_month = start_month

            while (
                current_year,
                current_month,
            ) <= (
                last_year,
                last_month,
            ):

                generation_key = (
                    f"monthly-{current_year}-{current_month:02d}"
                )

                if not ReportHistory.objects.filter(
                    generation_key=generation_key,
                ).exists():

                    reports.append(
                        ReportsService.create_monthly_report(
                            year=current_year,
                            month=current_month,
                        )
                    )

                # ماه بعدی
                if current_month == 12:
                    current_year += 1
                    current_month = 1
                else:
                    current_month += 1

        # ============================================================
        # YEARLY REPORTS
        # ============================================================

        if yearly:

            # آخرین سال کاملاً تکمیل‌شده
            last_year = year - 1

            # قدیمی‌ترین گزارش سالانه موجود
            first_yearly = (
                ReportHistory.objects
                .filter(
                    report_type="yearly",
                    period_year__isnull=False,
                )
                .order_by(
                    "period_year",
                )
                .first()
            )

            if first_yearly:
                start_year = first_yearly.period_year
            else:
                start_year = last_year

            for current_year in range(
                start_year,
                last_year + 1,
            ):

                generation_key = (
                    f"yearly-{current_year}"
                )

                if not ReportHistory.objects.filter(
                    generation_key=generation_key,
                ).exists():

                    reports.append(
                        ReportsService.create_yearly_report(
                            year=current_year,
                        )
                    )

        return reports

    @staticmethod
    @transaction.atomic
    def register_backup(
        file_name,
        file_path,
        size,
        backup_type="automatic",
        status="success",
        description="",
        created_by=None,
    ):

        return BackupHistory.objects.create(

            file_name=file_name,

            file_path=file_path,

            size=size,

            backup_type=backup_type,

            status=status,

            description=description,

            created_by=created_by,
        )

    @staticmethod
    def get_backup_history():

        return BackupHistory.objects.all()

    @staticmethod
    def get_last_backup():

        return BackupHistory.objects.order_by(
            "-created_at"
        ).first()

    @staticmethod
    def get_successful_backups():

        return BackupHistory.objects.filter(
            status="success"
        )

    @staticmethod
    def get_failed_backups():

        return BackupHistory.objects.filter(
            status="failed"
        )

    @staticmethod
    def delete_backup(backup_id):

        backup = BackupHistory.objects.filter(
            id=backup_id
        ).first()

        if not backup:
            return False

        backup.delete()

        return True

    @staticmethod
    def run_monthly_report():

        return ReportsService.generate_scheduled_reports(
            monthly=True,
            yearly=False,
        )

    @staticmethod
    def run_yearly_report():

        return ReportsService.generate_scheduled_reports(
            monthly=False,
            yearly=True,
        )

    @staticmethod
    def export_management_pdf():

        report = ReportsService.get_last_report()

        if not report:
            return b""

        file_object = ReportsService._get_report_file_object(
            report
        )

        try:
            return file_object.read()
        finally:
            file_object.close()

    @staticmethod
    def export_complete_pdf(
        start_date=None,
        end_date=None,
    ):

        report_data = ReportsService.generate_complete_report(
            start_date=start_date,
            end_date=end_date,
        )

        return ManagementReportPDFGenerator().generate_monthly_report({
            **report_data,
            "report_type": "custom",
            "month_label": "گزارش سفارشی",
            "period_start": report_data.get("period_start"),
            "period_end": report_data.get("period_end"),
            "ml": report_data.get("ml", {}),
        })

    @staticmethod
    def send_monthly_report_to_manager():

        from .tasks import ScheduledTasks

        return ScheduledTasks.send_monthly_report_to_manager()

    @staticmethod
    def send_yearly_report_to_manager():

        from .tasks import ScheduledTasks

        return ScheduledTasks.send_yearly_report_to_manager()

    @staticmethod
    def run_all_scheduled_tasks():

        return {

            "monthly_report":
                ReportsService.run_monthly_report(),

            "yearly_report":
                ReportsService.run_yearly_report(),

        }

    @staticmethod
    def get_report_history():

        return ReportHistory.objects.all()

    @staticmethod
    def get_last_report():

        return ReportHistory.objects.order_by(
            "-created_at"
        ).first()

    @staticmethod
    def get_reports_by_type(report_type):

        return ReportHistory.objects.filter(
            report_type=report_type
        )

    @staticmethod
    def get_reports_by_format(report_format):

        return ReportHistory.objects.filter(
            report_format=report_format
        )

    @staticmethod
    def delete_report(report_id):

        return False

    @staticmethod
    def export_excel(
        report_data,
        sheet_name="گزارش",
    ):

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = sheet_name[:31]

        row = 1

        for key, value in report_data.items():

            sheet.cell(
                row=row,
                column=1,
                value=str(key),
            )

            if isinstance(value, (dict, list)):
                value = str(value)

            sheet.cell(
                row=row,
                column=2,
                value=value,
            )

            row += 1

        output = BytesIO()
        workbook.save(output)

        return output.getvalue()

    @staticmethod
    def export_csv(report_data):

        output = StringIO()

        writer = csv.writer(
            output
        )

        writer.writerow([
            "شاخص",
            "مقدار",
        ])

        for key, value in report_data.items():

            writer.writerow([
                key,
                value,
            ])

        return output.getvalue()
